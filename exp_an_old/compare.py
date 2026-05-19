"""
Thu thập val_metrics.json từ các output dirs và xuất Excel so sánh.

Usage:
  python compare.py --dirs outputs/exp1_* --out results/exp1_compare.xlsx
  python compare.py --dirs outputs/exp1_viscam_binary outputs/exp1_vanilla_binary --out exp1.xlsx
"""

import argparse
import json
import os
import glob

import pandas as pd


METRIC_COLS = [
    ("val_loss",             "Val Loss"),
    ("dialogue_accuracy",    "Accuracy"),
    ("dialogue_f1",          "F1"),
    ("auroc",                "AUROC"),
    ("detection_rate",       "Detection Rate"),
    ("avg_detection_delay",  "Avg Alert Turn"),
    ("false_alarm_rate",     "False Alarm Rate"),
    ("median_alert_turn",    "Median Alert Turn"),
    ("median_lead_frac",     "Median Lead Frac"),
    ("alert_at_half",        "Alert in 1st Half"),
    ("num_detected",         "TP"),
    ("num_scam",             "# Scam"),
    ("num_false_alarms",     "FP"),
    ("num_harmless",         "# Harmless"),
]


def load_run(output_dir: str) -> dict | None:
    metrics_path = os.path.join(output_dir, "best_model", "val_metrics.json")
    config_path  = os.path.join(output_dir, "best_model", "config.json")

    if not os.path.exists(metrics_path):
        print(f"  [SKIP] No val_metrics.json in {output_dir}")
        return None

    with open(metrics_path) as f:
        metrics = json.load(f)

    cfg = {}
    if os.path.exists(config_path):
        with open(config_path) as f:
            cfg = json.load(f)

    run_name = os.path.basename(output_dir)
    row = {
        "Run":              run_name,
        "Mode":             cfg.get("mode", "binary"),
        "Model":            cfg.get("model_name", "").split("/")[-1],
        "w_max":            cfg.get("w_max", ""),
        "w_floor":          cfg.get("w_floor", ""),
        "patience_weight":  cfg.get("patience_weight", ""),
        "class_w_harmless": cfg.get("class_weight_harmless", ""),
        "lr":               cfg.get("lr", ""),
        "batch_size":       cfg.get("batch_size", ""),
        "grad_accum":       cfg.get("grad_accum_steps", ""),
    }
    for key, label in METRIC_COLS:
        row[label] = metrics.get(key, "")

    return row


def build_excel(rows: list[dict], out_path: str):
    os.makedirs(os.path.dirname(out_path) or ".", exist_ok=True)
    df = pd.DataFrame(rows)

    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Results")

        ws  = writer.sheets["Results"]
        wb  = writer.book

        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
        from openpyxl.formatting.rule import ColorScaleRule

        # Header style
        header_fill = PatternFill("solid", fgColor="1F3864")
        header_font = Font(bold=True, color="FFFFFF", size=11)
        thin = Side(style="thin", color="CCCCCC")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)

        for cell in ws[1]:
            cell.fill      = header_fill
            cell.font      = header_font
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
            cell.border    = border

        # Alternating row colors
        light = PatternFill("solid", fgColor="EEF2FF")
        for row_idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
            fill = light if row_idx % 2 == 0 else None
            for cell in row:
                cell.border    = border
                cell.alignment = Alignment(horizontal="center")
                if fill:
                    cell.fill = fill

        # Auto column width
        for col_idx, col in enumerate(ws.columns, 1):
            max_len = max((len(str(c.value or "")) for c in col), default=8)
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 3, 28)

        # Color scale cho F1 (higher = greener)
        metric_headers = [label for _, label in METRIC_COLS]
        for col_idx, cell in enumerate(ws[1], 1):
            if cell.value in ("F1", "AUROC", "Detection Rate", "Alert in 1st Half"):
                col_letter = get_column_letter(col_idx)
                ws.conditional_formatting.add(
                    f"{col_letter}2:{col_letter}{ws.max_row}",
                    ColorScaleRule(
                        start_type="min", start_color="FFB3B3",
                        end_type="max",   end_color="B3FFB3",
                    ),
                )
            elif cell.value in ("Val Loss", "False Alarm Rate", "Avg Alert Turn"):
                col_letter = get_column_letter(col_idx)
                ws.conditional_formatting.add(
                    f"{col_letter}2:{col_letter}{ws.max_row}",
                    ColorScaleRule(
                        start_type="min", start_color="B3FFB3",
                        end_type="max",   end_color="FFB3B3",
                    ),
                )

        ws.freeze_panes = "A2"

    print(f"Saved: {out_path}  ({len(rows)} runs)")


def main():
    parser = argparse.ArgumentParser(description="Compare experiment results → Excel")
    parser.add_argument("--dirs", nargs="+", required=True,
                        help="Output dirs (glob OK, e.g. outputs/exp1_*)")
    parser.add_argument("--out",  default="results/compare.xlsx",
                        help="Output Excel path")
    args = parser.parse_args()

    # Expand globs
    dirs = []
    for pattern in args.dirs:
        expanded = sorted(glob.glob(pattern))
        dirs.extend(expanded if expanded else [pattern])

    rows = []
    for d in dirs:
        row = load_run(d)
        if row:
            rows.append(row)

    if not rows:
        print("No completed runs found.")
        return

    build_excel(rows, args.out)


if __name__ == "__main__":
    main()
